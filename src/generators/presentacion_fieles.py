#!/usr/bin/env python3
"""
Generador de presentaciones PPTX para la asamblea (sin acordes)
Formato 16:9 basado en 274 Domingo 21 06 2026.pptx
"""

import os
import logging
import re
import sqlite3
from pathlib import Path
from typing import Dict, List, Optional, Tuple, Union

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

# Configurar logging
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s"
)
logger = logging.getLogger(__name__)

# Constantes
DEFAULT_PROJECT_DIR = Path.home() / "proyectos" / "CCE-M5-Web-Presentaciones"
PROJECT_DIR = Path(os.environ.get("CCE_PROJECT_DIR", DEFAULT_PROJECT_DIR))
DB_PATH = PROJECT_DIR / "data" / "db.sqlite3"
OUTPUT_DIR = PROJECT_DIR / "presentaciones"
TEMPLATE_PATH = PROJECT_DIR / "data" / "templates" / "274_Domingo_21_06_2026.pptx"

# Colores litúrgicos
COLORES_LITURGICOS = {
    "verde": RGBColor(0x4C, 0xAF, 0x50),
    "violeta": RGBColor(0x9C, 0x27, 0xB0),
    "blanco": RGBColor(0xFF, 0xFF, 0xFF),
    "rojo": RGBColor(0xF4, 0x43, 0x36),
    "rosa": RGBColor(0xE9, 0x1E, 0x63),
    "negro": RGBColor(0x21, 0x21, 0x21),
}

COLOR_TEXTO_CLARO = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_TEXTO_OSCURO = RGBColor(0x33, 0x33, 0x33)
COLOR_CITA = RGBColor(0x66, 0x66, 0x66)
COLOR_MOMENTO = RGBColor(0x99, 0x99, 0x99)

# Tipografía y márgenes
FUENTE_PRINCIPAL = "Calibri"
MARGEN_LATERAL = Inches(0.5)
MARGEN_SUPERIOR = Inches(0.3)
ANCHO_UTIL = Inches(12.333)

# Momentos del culto en orden
MOMENTOS: List[str] = [
    "entrada", "perdon", "gloria", "salmo", "aleluya",
    "ofertorio", "santo", "padre_nuestro", "paz",
    "comunion", "maria", "despedida",
]


def _color_para_texto(color_fondo: RGBColor) -> RGBColor:
    """Devuelve blanco para fondos oscuros, negro/gris para fondos claros."""
    # RGBColor es una namedtuple-like: accedemos por índice.
    r, g, b = color_fondo[0], color_fondo[1], color_fondo[2]
    luminancia = (0.299 * r + 0.587 * g + 0.114 * b) / 255
    return COLOR_TEXTO_CLARO if luminancia < 0.5 else COLOR_TEXTO_OSCURO


def _sanitizar_nombre_archivo(nombre: str) -> str:
    """Elimina caracteres peligrosos del nombre de archivo para evitar path traversal."""
    nombre_limpio = re.sub(r"[<>:\"|?*\x00-\x1f]", "_", nombre)
    nombre_limpio = nombre_limpio.replace("/", "_").replace("\\", "_")
    nombre_limpio = nombre_limpio.strip(". ")
    return nombre_limpio or "presentacion"


def _sanitizar_xml(texto: Optional[str]) -> str:
    """Elimina caracteres de control ilegales en XML (0x00-0x08, 0x0B-0x0C, 0x0E-0x1F)."""
    if not texto:
        return ""
    return re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F]", "", str(texto))


def _aplicar_fuente_y_margen(text_frame, color: RGBColor, font_size: Pt,
                              bold: bool = False, italic: bool = False,
                              alignment: PP_ALIGN = PP_ALIGN.LEFT,
                              line_spacing: Optional[float] = None) -> None:
    """Aplica fuente, tamaño, color, estilo y alineación a un párrafo."""
    p = text_frame.paragraphs[0]
    p.font.name = FUENTE_PRINCIPAL
    p.font.size = font_size
    p.font.bold = bold
    p.font.italic = italic
    p.font.color.rgb = color
    p.alignment = alignment
    if line_spacing is not None:
        p.line_spacing = line_spacing


class GeneradorPPTX:
    """Generador de presentaciones litúrgicas en formato 16:9."""

    def __init__(
        self,
        db_path: Optional[Union[str, Path]] = None,
        template_path: Optional[Union[str, Path]] = None,
        output_dir: Optional[Union[str, Path]] = None,
    ) -> None:
        self.db_path = Path(db_path) if db_path else DB_PATH
        self.template_path = Path(template_path) if template_path else TEMPLATE_PATH
        self.output_dir = Path(output_dir) if output_dir else OUTPUT_DIR

        if not self.template_path.is_file():
            raise FileNotFoundError(
                f"No se encontró el template PPTX: {self.template_path}"
            )

        self.prs = Presentation(str(self.template_path))
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def _get_connection(self) -> sqlite3.Connection:
        """Obtiene conexión a la base de datos con filas accesibles por nombre."""
        conn = sqlite3.connect(str(self.db_path))
        conn.row_factory = sqlite3.Row
        return conn

    def _crear_textbox(
        self,
        slide,
        left: float,
        top: float,
        width: float,
        height: float,
        texto: str,
        color: RGBColor,
        font_size: Pt,
        bold: bool = False,
        italic: bool = False,
        alignment: PP_ALIGN = PP_ALIGN.LEFT,
        line_spacing: Optional[float] = None,
        word_wrap: bool = True,
    ) -> None:
        """Crea un textbox en la diapositiva y aplica estilos consistentes."""
        textbox = slide.shapes.add_textbox(
            Inches(left), Inches(top), Inches(width), Inches(height)
        )
        tf = textbox.text_frame
        tf.word_wrap = word_wrap
        _aplicar_fuente_y_margen(
            tf, color, font_size, bold=bold, italic=italic,
            alignment=alignment, line_spacing=line_spacing,
        )
        tf.paragraphs[0].text = _sanitizar_xml(texto)

    def _color_fondo(self, color_fondo: str) -> RGBColor:
        """Resuelve el color de fondo a partir de su nombre."""
        return COLORES_LITURGICOS.get(color_fondo, COLORES_LITURGICOS["verde"])

    def crear_diapositiva_titulo(
        self,
        titulo: str,
        subtitulo: str,
        color_fondo: str = "verde",
    ) -> object:
        """Crea diapositiva de título con color litúrgico y texto contrastado."""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        fondo = self._color_fondo(color_fondo)
        color_texto = _color_para_texto(fondo)

        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = fondo

        # Título principal
        self._crear_textbox(
            slide,
            left=1.0,
            top=2.5,
            width=11.333,
            height=1.5,
            texto=titulo,
            color=color_texto,
            font_size=Pt(54),
            bold=True,
            alignment=PP_ALIGN.CENTER,
        )

        # Subtítulo
        self._crear_textbox(
            slide,
            left=1.0,
            top=4.2,
            width=11.333,
            height=1.5,
            texto=subtitulo,
            color=color_texto,
            font_size=Pt(28),
            alignment=PP_ALIGN.CENTER,
        )

        return slide

    def crear_diapositiva_lectura(
        self,
        titulo: str,
        cita: str,
        texto: str,
        color_fondo: str = "blanco",
    ) -> object:
        """Crea diapositiva de lectura usando el color de fondo recibido."""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        fondo = self._color_fondo(color_fondo)
        color_texto = _color_para_texto(fondo)

        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = fondo

        # Título de la lectura
        self._crear_textbox(
            slide,
            left=0.5,
            top=0.3,
            width=12.333,
            height=0.6,
            texto=titulo,
            color=color_texto,
            font_size=Pt(32),
            bold=True,
        )

        # Cita
        self._crear_textbox(
            slide,
            left=0.5,
            top=1.0,
            width=12.333,
            height=0.6,
            texto=cita,
            color=COLOR_CITA if fondo == COLORES_LITURGICOS["blanco"] else color_texto,
            font_size=Pt(20),
            italic=True,
        )

        # Texto de la lectura
        texto_limpio = _sanitizar_xml(texto)
        if len(texto_limpio) > 500:
            texto_limpio = texto_limpio[:500] + "..."

        self._crear_textbox(
            slide,
            left=0.5,
            top=1.8,
            width=12.333,
            height=5.0,
            texto=texto_limpio,
            color=color_texto,
            font_size=Pt(18),
            line_spacing=1.5,
        )

        return slide

    def crear_diapositiva_cancion(
        self,
        titulo: str,
        letra: str,
        momento: str,
        color_fondo: str = "blanco",
    ) -> object:
        """Crea diapositiva de canción sin acordes usando el color de fondo recibido."""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        fondo = self._color_fondo(color_fondo)
        color_texto = _color_para_texto(fondo)

        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = fondo

        # Título del momento
        self._crear_textbox(
            slide,
            left=0.5,
            top=0.3,
            width=12.333,
            height=0.5,
            texto=momento.upper(),
            color=COLOR_MOMENTO if fondo == COLORES_LITURGICOS["blanco"] else color_texto,
            font_size=Pt(24),
            bold=True,
        )

        # Título de la canción
        self._crear_textbox(
            slide,
            left=0.5,
            top=0.9,
            width=12.333,
            height=0.5,
            texto=titulo,
            color=color_texto,
            font_size=Pt(36),
            bold=True,
        )

        # Letra limpia: quitar metadatos del blogspot y líneas de acordes sueltas.
        from src.parsers.acordes_parser_v3 import limpiar_texto_cancion, AcordesParser
        letra_limpia = limpiar_texto_cancion(letra)
        parser = AcordesParser()
        lineas_letra: list[str] = []
        for linea in parser.parsear_cancion_completa(letra_limpia):
            if linea.tipo == "letra" and linea.letra.strip():
                lineas_letra.append(linea.letra.strip())
            elif linea.tipo == "acordes_letra" and linea.letra.strip():
                lineas_letra.append(linea.letra.strip())
            elif linea.tipo == "sección" and linea.texto.strip():
                lineas_letra.append(f"[{linea.texto.strip().strip('[]')}]")
        letra_visual = "\n".join(lineas_letra)
        letra_visual = _sanitizar_xml(letra_visual)
        if len(letra_visual) > 1200:
            letra_visual = letra_visual[:1200] + "..."

        self._crear_textbox(
            slide,
            left=0.5,
            top=1.8,
            width=12.333,
            height=5.2,
            texto=letra_visual,
            color=color_texto,
            font_size=Pt(20),
            alignment=PP_ALIGN.CENTER,
            line_spacing=1.5,
        )

        return slide

    def _precargar_canciones(
        self,
        cursor: sqlite3.Cursor,
        canciones_ids: Dict[str, int],
    ) -> Dict[int, sqlite3.Row]:
        """Precarga todas las canciones necesarias con un único query IN(...)."""
        ids = [cid for cid in canciones_ids.values() if isinstance(cid, int) and cid > 0]
        if not ids:
            return {}

        placeholders = ",".join("?" * len(ids))
        cursor.execute(
            f"SELECT id, titulo, letra_sin_acordes FROM canciones WHERE id IN ({placeholders})",
            ids,
        )
        return {row["id"]: row for row in cursor.fetchall()}

    def generar_presentacion(
        self,
        fecha_domingo: str,
        lectura_id: int,
        canciones_ids: Dict[str, int],
    ) -> Optional[Path]:
        """Genera una presentación completa y devuelve la ruta del archivo generado."""
        logger.info("Generando presentación para: %s", fecha_domingo)

        try:
            with self._get_connection() as conn:
                cursor = conn.cursor()

                cursor.execute(
                    """
                    SELECT COALESCE(celebracion, domingo) AS celebracion,
                           temporada, color_liturgico,
                           primera_lectura_cita, primera_lectura_texto,
                           salmo_cita, salmo_antifona, salmo_texto,
                           segunda_lectura_cita, segunda_lectura_texto,
                           evangelio_cita, evangelio_texto
                    FROM lecturas WHERE id = ?
                    """,
                    (lectura_id,),
                )
                lectura = cursor.fetchone()
                if not lectura:
                    logger.error("No se encontró lectura con id=%s", lectura_id)
                    return None

                celebracion = lectura["celebracion"] or "Celebración del Domingo"
                color = lectura["color_liturgico"] or "verde"

                # Portada
                self.crear_diapositiva_titulo(
                    celebracion,
                    fecha_domingo,
                    color,
                )

                # Lecturas
                if lectura["primera_lectura_texto"]:
                    self.crear_diapositiva_lectura(
                        "Primera Lectura",
                        lectura["primera_lectura_cita"] or "",
                        lectura["primera_lectura_texto"],
                        color,
                    )

                if lectura["salmo_texto"]:
                    self.crear_diapositiva_lectura(
                        "Salmo Responsorial",
                        lectura["salmo_cita"] or "",
                        f"Antífona: {lectura['salmo_antifona'] or ''}\n\n{lectura['salmo_texto']}",
                        color,
                    )

                if lectura["segunda_lectura_texto"]:
                    self.crear_diapositiva_lectura(
                        "Segunda Lectura",
                        lectura["segunda_lectura_cita"] or "",
                        lectura["segunda_lectura_texto"],
                        color,
                    )

                if lectura["evangelio_texto"]:
                    self.crear_diapositiva_lectura(
                        "Evangelio",
                        lectura["evangelio_cita"] or "",
                        lectura["evangelio_texto"],
                        color,
                    )

                # Canciones
                canciones_cache = self._precargar_canciones(cursor, canciones_ids)

                for momento in MOMENTOS:
                    if momento not in canciones_ids:
                        continue

                    cancion_id = canciones_ids[momento]
                    cancion = canciones_cache.get(cancion_id)
                    if not cancion:
                        logger.warning("Canción no encontrada para '%s' (id=%s)", momento, cancion_id)
                        continue

                    self.crear_diapositiva_cancion(
                        cancion["titulo"] or "Canción",
                        cancion["letra_sin_acordes"] or "",
                        momento,
                        "blanco",
                    )

                # Guardar
                filename = f"{_sanitizar_nombre_archivo(fecha_domingo)}_presentacion.pptx"
                filepath = self.output_dir / filename
                self.prs.save(str(filepath))

                logger.info("Presentación generada: %s", filepath)
                logger.info("Diapositivas: %s", len(self.prs.slides))
                return filepath

        except sqlite3.Error as e:
            logger.error("Error de base de datos: %s", e)
            return None
        except FileNotFoundError as e:
            logger.error("Archivo no encontrado: %s", e)
            return None
        except PermissionError as e:
            logger.error("Permiso denegado al escribir PPTX: %s", e)
            return None
        except Exception as e:
            logger.error("Error al generar presentación: %s", e)
            return None


if __name__ == "__main__":
    raise SystemExit("Usa el script demo para probar el generador.")
